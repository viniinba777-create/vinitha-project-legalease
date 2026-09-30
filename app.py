import os
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from google import genai
from docx import Document
from fpdf import FPDF


# =========================================================
# LOAD ENVIRONMENT VARIABLES
# =========================================================

load_dotenv()

# Also check the main LegalEase folder for .env
project_env = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    ".env"
)

if os.path.exists(project_env):
    load_dotenv(project_env)


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="centered"
)


# =========================================================
# HEADER
# =========================================================

st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")

st.info(
    "LegalEase creates AI-generated document drafts. "
    "It is not a substitute for legal advice."
)


# =========================================================
# DOCUMENT TYPE
# =========================================================

document_type = st.selectbox(
    "Select Document Type",
    [
        "Rental Agreement",
        "Leave/Application Letter",
        "General Notice",
        "Affidavit-style Draft",
        "NDA",
        "Employment Agreement"
    ]
)


# =========================================================
# USER INFORMATION
# =========================================================

name = st.text_input("Name")

address = st.text_area("Address")

date = st.date_input("Date")

details = st.text_area(
    "Additional Details",
    placeholder="Enter the important details for the document..."
)


# =========================================================
# CREATE DOCX
# =========================================================

def create_docx(document_text, document_type):

    doc = Document()

    doc.add_heading(
        document_type,
        level=1
    )

    doc.add_paragraph(
        "AI-Generated Draft"
    )

    doc.add_paragraph(
        "This document is an AI-generated draft and "
        "should be reviewed by an appropriate legal "
        "professional before use."
    )

    for paragraph in document_text.split("\n"):

        if paragraph.strip():

            doc.add_paragraph(
                paragraph.strip()
            )

    output = BytesIO()

    doc.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# CREATE PDF
# =========================================================

def create_pdf(document_text, document_type):

    pdf = FPDF()

    pdf.add_page()

    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )

    # Title
    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
        0,
        10,
        document_type,
        align="C"
    )

    pdf.ln(5)

    # Notice
    pdf.set_font(
        "Helvetica",
        "I",
        10
    )

    pdf.multi_cell(
        0,
        6,
        "AI-Generated Draft - This document should be "
        "reviewed by an appropriate legal professional "
        "before use."
    )

    pdf.ln(5)

    # Document content
    pdf.set_font(
        "Helvetica",
        "",
        11
    )

    # Make text compatible with the PDF font
    safe_text = (
        document_text
        .encode("latin-1", "replace")
        .decode("latin-1")
    )

    for paragraph in safe_text.split("\n"):

        if paragraph.strip():

            pdf.multi_cell(
                0,
                7,
                paragraph.strip()
            )

            pdf.ln(2)

    return bytes(pdf.output())


# =========================================================
# GENERATE DOCUMENT
# =========================================================

if st.button("Generate Document"):

    # -----------------------------------------------------
    # API KEY
    # -----------------------------------------------------

    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:

        st.error(
            "Gemini API key is not configured. "
            "Please add GEMINI_API_KEY to your .env file."
        )

        st.stop()


    # -----------------------------------------------------
    # GEMINI CLIENT
    # -----------------------------------------------------

    client = genai.Client(
        api_key=api_key
    )


    # -----------------------------------------------------
    # PROMPT
    # -----------------------------------------------------

    prompt = f"""
You are an AI legal document drafting assistant.

Create a professional legal document DRAFT.

Document type:
{document_type}

Name:
{name}

Address:
{address}

Date:
{date}

Additional details:
{details}

Instructions:

1. Create a clear and professional legal document draft.

2. The document MUST match the selected document type:
{document_type}

3. Use only the information provided by the user.

4. Do not invent important personal, financial, or legal facts.

5. Organize the document with clear headings and sections.

6. Include the provided date.

7. Clearly identify the person or party provided.

8. Use formal language appropriate for the selected document type.

9. Include relevant clauses based on the details provided.

10. Clearly state that this is an AI-generated draft.

11. State that the document should be reviewed by an appropriate
legal professional before use.

12. Return only the document text.

This document is a draft and is not legal advice.
"""


    # -----------------------------------------------------
    # CALL GEMINI
    # -----------------------------------------------------

    try:

        response = client.interactions.create(
            model="gemini-3.5-flash-lite",
            input=prompt
        )

        document_text = response.output_text


        # -------------------------------------------------
        # DISPLAY GENERATED DOCUMENT
        # -------------------------------------------------

        st.subheader("Generated Document")

        st.text_area(
            "Document Draft",
            document_text,
            height=500
        )


        # -------------------------------------------------
        # FILE NAME
        # -------------------------------------------------

        safe_name = (
            document_type
            .replace(" ", "_")
            .replace("/", "_")
            .replace("\\", "_")
        )


        # =================================================
        # TXT
        # =================================================

        txt_file = document_text.encode("utf-8")

        st.download_button(
            label="📄 Download as TXT",
            data=txt_file,
            file_name=f"{safe_name}.txt",
            mime="text/plain"
        )


        # =================================================
        # DOCX
        # =================================================

        docx_file = create_docx(
            document_text,
            document_type
        )

        st.download_button(
            label="📝 Download as DOCX",
            data=docx_file,
            file_name=f"{safe_name}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )


        # =================================================
        # PDF
        # =================================================

        pdf_file = create_pdf(
            document_text,
            document_type
        )

        st.download_button(
            label="📕 Download as PDF",
            data=pdf_file,
            file_name=f"{safe_name}.pdf",
            mime="application/pdf"
        )


    except Exception as e:

        st.error(
            f"Error generating document: {e}"
        )